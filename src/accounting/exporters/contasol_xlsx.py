"""
Exportador para DELSOL Contasol (Libro Excel .xlsx)
Genera el libro estructurado compatible con la funcionalidad nativa de Contasol
'Importar asientos desde hoja de cálculo'.
Columnas: Diario, Fecha, Asiento, Subcuenta, Concepto, Documento, Debe, Haber, Contrapartida.
"""
import io
from pathlib import Path
from typing import List, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from src.accounting.double_entry import DoubleEntryJournal, normalize_pgc_subaccount


class ContasolXlsxExporter:
    HEADERS = [
        "Diario",
        "Fecha",
        "Asiento",
        "Subcuenta",
        "Concepto",
        "Documento",
        "Debe",
        "Haber",
        "Contrapartida",
    ]

    @staticmethod
    def export_xlsx(
        journals: List[DoubleEntryJournal],
        output_file_or_path: Union[str, Path, None] = None,
        account_length: int = 9,
    ) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Asientos Contables"

        # Estilo de cabecera
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        ws.append(ContasolXlsxExporter.HEADERS)
        for col_idx in range(1, len(ContasolXlsxExporter.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for journal in journals:
            if not journal.is_balanced:
                journal.check_balance()
                if not journal.is_balanced:
                    raise ValueError(
                        f"Asiento {journal.entry_number} descuadrado ({journal.imbalance_cents} céntimos). "
                        f"Contasol bloquearía el diario."
                    )

            date_str = journal.entry_date.strftime("%d/%m/%Y")
            entry_num = journal.entry_number

            for line in journal.lines:
                acct = normalize_pgc_subaccount(line.account_code, account_length)
                contra = normalize_pgc_subaccount(line.contrapartida or "", account_length)
                debit_val = round(line.debit_cents / 100.0, 2)
                credit_val = round(line.credit_cents / 100.0, 2)

                row_data = [
                    1,                 # Diario 1 (General)
                    date_str,          # Fecha
                    entry_num,         # Asiento
                    acct,              # Subcuenta
                    (line.concept or "")[:50], # Concepto
                    (line.document_number or "")[:15], # Documento
                    debit_val,         # Debe
                    credit_val,        # Haber
                    contra,            # Contrapartida
                ]
                ws.append(row_data)

        # Autoajuste de anchos de columna
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        if output_file_or_path:
            wb.save(output_file_or_path)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
