"""
Generador de Datos de Prueba: Carpeta 'prueba Data' con 20 Facturas Reales Inventadas
Genera:
  1. Cartera consolidada en Excel (.xlsx) con estilos profesionales
  2. Cartera en CSV (.csv) con formato español (; y UTF-8-BOM)
  3. Cartera en JSON (.json)
  4. 20 Facturas individuales en PDF con maquetación ejecutiva (ReportLab)
  5. Fichero bancario oficial Norma 43 (.n43) sincronizado para conciliación inmediata
  6. Fichero de pasarela Stripe (.csv) para prueba de lotes y comisiones
  7. Guía rápida de uso
"""
import os
import sys
import json
import csv
import re
from datetime import datetime, date, timedelta
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

BASE_DIR = Path("E:/Games/Proyecto")
TARGET_DIR = BASE_DIR / "prueba Data"
PDF_DIR = TARGET_DIR / "facturas_pdf"

TARGET_DIR.mkdir(parents=True, exist_ok=True)
PDF_DIR.mkdir(parents=True, exist_ok=True)

# 20 Facturas B2B / E-Commerce realistas para España
INVOICES = [
    {
        "num": "FRA-2026-001",
        "client": "Logística Exprés Ibérica S.L.",
        "cif": "B-87654321",
        "address": "Avda. de la Industria 45, Pol. Ind. Alcobendas",
        "city": "28108 Alcobendas (Madrid)",
        "subcuenta": "430000001",
        "date": "2026-10-02",
        "due": "2026-10-02",
        "base": 1033.06,
        "iva_rate": 21,
        "iva": 216.94,
        "total": 1250.00,
        "concept": "Servicios de transporte y distribución paletizada Q3",
        "qty": 1,
        "status": "Cobrada por Stripe",
    },
    {
        "num": "FRA-2026-002",
        "client": "Telefónica Soluciones Digitales S.A.",
        "cif": "A-28001234",
        "address": "Ronda de la Comunicación s/n, Edificio Central",
        "city": "28050 Madrid",
        "subcuenta": "430000002",
        "date": "2026-10-03",
        "due": "2026-10-03",
        "base": 735.95,
        "iva_rate": 21,
        "iva": 154.55,
        "total": 890.50,
        "concept": "Licencia software cloud computing y soporte 24/7",
        "qty": 1,
        "status": "Cobrada por Stripe",
    },
    {
        "num": "FRA-2026-003",
        "client": "Almacenes Centrales Madrid S.L.",
        "cif": "B-81234567",
        "address": "Calle Resina 12, Pol. Ind. Villaverde",
        "city": "28021 Madrid",
        "subcuenta": "430000003",
        "date": "2026-10-05",
        "due": "2026-10-05",
        "base": 2024.79,
        "iva_rate": 21,
        "iva": 425.21,
        "total": 2450.00,
        "concept": "Suministro estanterías industriales modulares tipo rack",
        "qty": 2,
        "status": "Cobrada por Transferencia B2B",
    },
    {
        "num": "FRA-2026-004",
        "client": "Distribuciones Alimentarias del Sur S.L.",
        "cif": "B-41889900",
        "address": "Ctra. de Carmona km 3.5",
        "city": "41008 Sevilla",
        "subcuenta": "430000004",
        "date": "2026-10-05",
        "due": "2026-10-05",
        "base": 574.38,
        "iva_rate": 21,
        "iva": 120.62,
        "total": 695.00,
        "concept": "Lote embalajes isotérmicos para alimentación fresca",
        "qty": 5,
        "status": "Cobrada por Stripe",
    },
    {
        "num": "FRA-2026-005",
        "client": "Consultoría Tecnológica Novatech S.L.",
        "cif": "B-91827364",
        "address": "Paseo de la Castellana 140, Planta 8",
        "city": "28046 Madrid",
        "subcuenta": "430000005",
        "date": "2026-10-06",
        "due": "2026-10-10",
        "base": 2561.98,
        "iva_rate": 21,
        "iva": 538.02,
        "total": 3100.00,
        "concept": "Auditoría de ciberseguridad y test de intrusión web",
        "qty": 1,
        "status": "Pendiente de cobro",
    },
    {
        "num": "FRA-2026-006",
        "client": "Construcciones & Reformas Mediterráneo S.L.",
        "cif": "B-54892103",
        "address": "Calle Mayor 88, Pol. Ind. Pla de la Vallonga",
        "city": "03006 Alicante",
        "subcuenta": "430000006",
        "date": "2026-10-07",
        "due": "2026-10-12",
        "base": 1471.28,
        "iva_rate": 21,
        "iva": 308.97,
        "total": 1780.25,
        "concept": "Materiales aislantes térmicos y placas yeso laminado",
        "qty": 1,
        "status": "Cobrada por Redsys TPV",
    },
    {
        "num": "FRA-2026-007",
        "client": "Papelería e Imprenta Gráficas Cantabria S.L.",
        "cif": "B-33445566",
        "address": "Avda. de los Castros 23",
        "city": "39005 Santander",
        "subcuenta": "430000007",
        "date": "2026-10-07",
        "due": "2026-10-12",
        "base": 347.11,
        "iva_rate": 21,
        "iva": 72.89,
        "total": 420.00,
        "concept": "Impresión de catálogos corporativos y flyers A5",
        "qty": 1000,
        "status": "Cobrada por Redsys TPV",
    },
    {
        "num": "FRA-2026-008",
        "client": "CONSTRUCCIONES HNOS. MARTINEZ S.L.",
        "cif": "B-41982737",
        "address": "Calle Torneo 72, Nave 4",
        "city": "41002 Sevilla",
        "subcuenta": "430000008",
        "date": "2026-10-08",
        "due": "2026-10-10",
        "base": 1198.35,
        "iva_rate": 21,
        "iva": 251.65,
        "total": 1450.00,
        "concept": "Instalación cuadros eléctricos trifásicos y canalización",
        "qty": 1,
        "status": "Cobrada por Transferencia B2B",
    },
    {
        "num": "FRA-2026-009",
        "client": "AMAZON SERVICES EUROPE S.A.R.L.",
        "cif": "W0185698H",
        "address": "38 Avenue John F. Kennedy",
        "city": "L-1855 Luxembourg",
        "subcuenta": "430000009",
        "date": "2026-10-08",
        "due": "2026-10-15",
        "base": 2410.00,
        "iva_rate": 0,
        "iva": 0.00,
        "total": 2410.00,
        "concept": "Servicios de logística e intermediación comercial europea (ISP)",
        "qty": 1,
        "status": "Cobrada por Transferencia B2B",
    },
    {
        "num": "FRA-2026-010",
        "client": "Asesoría Fiscal & Laboral Castelar S.L.P.",
        "cif": "B-86543219",
        "address": "Calle Castelar 15, 2º Derecha",
        "city": "46001 Valencia",
        "subcuenta": "430000010",
        "date": "2026-10-09",
        "due": "2026-10-20",
        "base": 785.12,
        "iva_rate": 21,
        "iva": 164.88,
        "total": 950.00,
        "concept": "Honorarios consultoría fiscal cierre 3T e IS anticipado",
        "qty": 1,
        "status": "Pendiente de cobro",
    },
    {
        "num": "FRA-2026-011",
        "client": "Maquinaria Industrial Aragón S.L.",
        "cif": "B-50987654",
        "address": "Pol. Ind. Malpica, Calle E, Parcela 12",
        "city": "50016 Zaragoza",
        "subcuenta": "430000011",
        "date": "2026-10-09",
        "due": "2026-10-25",
        "base": 3983.47,
        "iva_rate": 21,
        "iva": 836.53,
        "total": 4820.00,
        "concept": "Mantenimiento preventivo y sustitución bombas hidráulicas",
        "qty": 2,
        "status": "Pendiente de cobro",
    },
    {
        "num": "FRA-2026-012",
        "client": "Farmacéutica & Cosmética Cataluña S.A.",
        "cif": "A-08123987",
        "address": "Avda. Diagonal 640",
        "city": "08017 Barcelona",
        "subcuenta": "430000012",
        "date": "2026-10-10",
        "due": "2026-10-10",
        "base": 1351.65,
        "iva_rate": 21,
        "iva": 283.85,
        "total": 1635.50,
        "concept": "Envases estériles de vidrio grado farmacopea (100ml)",
        "qty": 500,
        "status": "Cobrada por Stripe",
    },
    {
        "num": "FRA-2026-013",
        "client": "Hotelería & Alojamientos Turísticos Levante S.L.",
        "cif": "B-98765432",
        "address": "Paseo Marítimo 112",
        "city": "46730 Gandía (Valencia)",
        "subcuenta": "430000013",
        "date": "2026-10-10",
        "due": "2026-10-12",
        "base": 644.63,
        "iva_rate": 21,
        "iva": 135.37,
        "total": 780.00,
        "concept": "Dotación lencería de cama y baño 100% algodón egipcio",
        "qty": 50,
        "status": "Cobrada por Redsys TPV",
    },
    {
        "num": "FRA-2026-014",
        "client": "Cítricos & Frutas del Turia S.L.",
        "cif": "B-46123456",
        "address": "Partida Vera s/n",
        "city": "46022 Valencia",
        "subcuenta": "430000014",
        "date": "2026-10-11",
        "due": "2026-10-25",
        "base": 1776.86,
        "iva_rate": 21,
        "iva": 373.14,
        "total": 2150.00,
        "concept": "Cajas de madera tratada para exportación agrícola",
        "qty": 300,
        "status": "Pendiente de cobro",
    },
    {
        "num": "FRA-2026-015",
        "client": "Frigoríficos & Pesca del Atlántico S.L.",
        "cif": "B-36987123",
        "address": "Muelle de Bouzas s/n",
        "city": "36208 Vigo (Pontevedra)",
        "subcuenta": "430000015",
        "date": "2026-10-11",
        "due": "2026-10-15",
        "base": 2809.92,
        "iva_rate": 21,
        "iva": 590.08,
        "total": 3400.00,
        "concept": "Reparación compresores de congelación rápida túnel B",
        "qty": 1,
        "status": "Cobrada por Transferencia B2B",
    },
    {
        "num": "FRA-2026-016",
        "client": "Textil y Confección Galicia S.L.",
        "cif": "B-15432987",
        "address": "Pol. Ind. de Sabón, Parcela 14",
        "city": "15142 Arteixo (A Coruña)",
        "subcuenta": "430000016",
        "date": "2026-10-12",
        "due": "2026-10-12",
        "base": 926.24,
        "iva_rate": 21,
        "iva": 194.51,
        "total": 1120.75,
        "concept": "Tejido técnico transpirable impermeable rollo 50m",
        "qty": 4,
        "status": "Cobrada por Stripe",
    },
    {
        "num": "FRA-2026-017",
        "client": "Metalúrgica del Besaya S.L.",
        "cif": "B-39876541",
        "address": "Pol. Ind. Barros, Calle 3",
        "city": "39400 Los Corrales de Buelna (Cantabria)",
        "subcuenta": "430000017",
        "date": "2026-10-12",
        "due": "2026-10-30",
        "base": 4628.10,
        "iva_rate": 21,
        "iva": 971.90,
        "total": 5600.00,
        "concept": "Perfilería de acero inoxidable grado marino AISI 316",
        "qty": 10,
        "status": "Pendiente de cobro",
    },
    {
        "num": "FRA-2026-018",
        "client": "Bodegas y Viñedos de La Rioja S.L.",
        "cif": "B-26543210",
        "address": "Camino de las Bodegas s/n",
        "city": "26006 Logroño (La Rioja)",
        "subcuenta": "430000018",
        "date": "2026-10-13",
        "due": "2026-10-20",
        "base": 735.54,
        "iva_rate": 21,
        "iva": 154.46,
        "total": 890.00,
        "concept": "Barricas de roble francés 225L para crianza enológica",
        "qty": 2,
        "status": "Cobrada por Transferencia B2B",
    },
    {
        "num": "FRA-2026-019",
        "client": "Embutidos & Jamones de Burgos S.L.",
        "cif": "B-09876543",
        "address": "Ctra. Madrid-Irún km 242",
        "city": "09007 Burgos",
        "subcuenta": "430000019",
        "date": "2026-10-13",
        "due": "2026-10-13",
        "base": 1107.44,
        "iva_rate": 21,
        "iva": 232.56,
        "total": 1340.00,
        "concept": "Bolsas de envasado al vacío retráctil alta barrera",
        "qty": 2000,
        "status": "Cobrada por Stripe",
    },
    {
        "num": "FRA-2026-020",
        "client": "Software & Soluciones Cloud Valladolid S.L.",
        "cif": "B-47123985",
        "address": "Parque Tecnológico de Boecillo, Parcela 2",
        "city": "47151 Boecillo (Valladolid)",
        "subcuenta": "430000020",
        "date": "2026-10-14",
        "due": "2026-10-31",
        "base": 2355.37,
        "iva_rate": 21,
        "iva": 494.63,
        "total": 2850.00,
        "concept": "Desarrollo API microservicios y despliegue contenedor Docker",
        "qty": 1,
        "status": "Pendiente de cobro",
    },
]

print(f"Total Facturas: {len(INVOICES)}")
total_suma = sum(inv["total"] for inv in INVOICES)
print(f"Suma Total Cartera: {total_suma:.2f} €")

# ==============================================================================
# 1. EXCEL (.xlsx) ESTILIZADO
# ==============================================================================
excel_path = TARGET_DIR / "cartera_facturas_emitidas_20.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Cartera Facturas Emitidas"

# Paleta Zinc/Fintech
header_fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
zebra_fill = PatternFill(start_color="F4F4F5", end_color="F4F4F5", fill_type="solid")
regular_font = Font(name="Segoe UI", size=9)
bold_font = Font(name="Segoe UI", size=9, bold=True)
mono_font = Font(name="Consolas", size=9)

thin_border = Border(
    left=Side(style="thin", color="E4E4E7"),
    right=Side(style="thin", color="E4E4E7"),
    top=Side(style="thin", color="E4E4E7"),
    bottom=Side(style="thin", color="E4E4E7")
)

headers = [
    "Nº Factura",
    "Cliente / Razón Social",
    "CIF / NIF",
    "Fecha Emisión",
    "Fecha Vencimiento",
    "Base Imponible (€)",
    "% IVA",
    "Cuota IVA (€)",
    "Total Factura (€)",
    "Subcuenta 430",
    "Estado de Cobro",
    "Concepto Principal"
]

ws.append(headers)
for col_num, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center" if "Fecha" in h or "Nº" in h or "CIF" in h else "left", vertical="center")
ws.row_dimensions[1].height = 24

for row_idx, inv in enumerate(INVOICES, start=2):
    row_data = [
        inv["num"],
        inv["client"],
        inv["cif"],
        inv["date"],
        inv["due"],
        inv["base"],
        f"{inv['iva_rate']}%",
        inv["iva"],
        inv["total"],
        inv["subcuenta"],
        inv["status"],
        inv["concept"]
    ]
    ws.append(row_data)
    is_even = (row_idx % 2 == 0)
    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=row_idx, column=col_num)
        c.font = regular_font
        c.border = thin_border
        if is_even:
            c.fill = zebra_fill
        
        # Formatos específicos
        if col_num in (6, 8, 9):
            c.number_format = "#,##0.00 €"
            c.alignment = Alignment(horizontal="right")
        elif col_num in (1, 3, 4, 5, 7, 10):
            c.alignment = Alignment(horizontal="center")
            if col_num in (1, 3, 10):
                c.font = mono_font

# Fila de Totales
total_row = len(INVOICES) + 2
ws.cell(row=total_row, column=2, value="TOTAL CARTERA").font = bold_font
ws.cell(row=total_row, column=6, value=f"=SUM(F2:F{total_row-1})").font = bold_font
ws.cell(row=total_row, column=6).number_format = "#,##0.00 €"
ws.cell(row=total_row, column=8, value=f"=SUM(H2:H{total_row-1})").font = bold_font
ws.cell(row=total_row, column=8).number_format = "#,##0.00 €"
ws.cell(row=total_row, column=9, value=f"=SUM(I2:I{total_row-1})").font = bold_font
ws.cell(row=total_row, column=9).number_format = "#,##0.00 €"

# Auto-ajustar anchos
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(excel_path)
print(f"Excel guardado en: {excel_path}")

# ==============================================================================
# 2. CSV (.csv) EN FORMATO ESPAÑOL (delimitador ';' y UTF-8-sig)
# ==============================================================================
csv_path = TARGET_DIR / "cartera_facturas_emitidas_20.csv"
with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f, delimiter=";")
    writer.writerow(headers)
    for inv in INVOICES:
        writer.writerow([
            inv["num"],
            inv["client"],
            inv["cif"],
            inv["date"],
            inv["due"],
            f"{inv['base']:.2f}".replace(".", ","),
            f"{inv['iva_rate']}%",
            f"{inv['iva']:.2f}".replace(".", ","),
            f"{inv['total']:.2f}".replace(".", ","),
            inv["subcuenta"],
            inv["status"],
            inv["concept"]
        ])
print(f"CSV guardado en: {csv_path}")

# ==============================================================================
# 3. JSON (.json) PARA IMPORTACIÓN DIRECTA DE CARTERA
# ==============================================================================
json_path = TARGET_DIR / "cartera_facturas_emitidas_20.json"
json_data = []
for inv in INVOICES:
    json_data.append({
        "invoice_number": inv["num"],
        "customer_name": inv["client"],
        "customer_fiscal_id": inv["cif"],
        "gross_amount_cents": int(round(inv["total"] * 100)),
        "gross_amount_eur": inv["total"],
        "issue_date": inv["date"],
        "due_date": inv["due"],
        "account_code": inv["subcuenta"],
        "concept": inv["concept"],
        "status": inv["status"]
    })
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(json_data, f, ensure_ascii=False, indent=2)
print(f"JSON guardado en: {json_path}")

# ==============================================================================
# 4. FICHERO BANCARIO NORMA 43 (.n43) OFICIAL PARA CONCILIACIÓN
# ==============================================================================
# Genera un extracto bancario con cobros correspondientes a las facturas
# FRA-2026-003 (2.450€), FRA-2026-008 (1.450€), FRA-2026-009 (2.410€), FRA-2026-015 (3.400€), FRA-2026-018 (890€)
n43_path = TARGET_DIR / "extracto_bancario_norma43.n43"
def pad(val, length, fill=' ', align='<'):
    s = str(val)[:length]
    if align == '<':
        return s.ljust(length, fill)
    else:
        return s.rjust(length, fill)

lines_n43 = []
# Cabecera 11 (80 caracteres exactos)
# 11 + 0049 (4) + 1500 (4) + 0012345678 (10) + 261001 (6) + 261015 (6) + 2 (1) + 00000001500000 (14) + 978 (3) + 3 (1) + 26 chars nombre + 3 espacios
h11 = "110049150000123456782610012610152000000015000009783FINDAFINITY DEMO SL          "
lines_n43.append(h11)

cobros_n43 = [
    {"date": "261005", "val": "261005", "cents": 245000, "doc": "0000000001", "ref1": "TRF ALMACENES CENTRALES", "c2": "TRANSF ALMACENES CENTRALES FRA-2026-003 CIF B81234567"},
    {"date": "261010", "val": "261010", "cents": 145000, "doc": "0000000002", "ref1": "TRF CONSTR HNOS MARTINEZ", "c2": "TRANSF CONSTR HNOS MARTINEZ FRA-2026-008 CIF B41982737"},
    {"date": "261015", "val": "261015", "cents": 241000, "doc": "0000000003", "ref1": "TRF AMAZON SERVICES EURO", "c2": "TRANSF AMAZON SERVICES EUROPE FRA-2026-009 NIF W0185698H"},
    {"date": "261015", "val": "261015", "cents": 340000, "doc": "0000000004", "ref1": "TRF FRIGORIFICOS ATLANTI", "c2": "TRANSF FRIGORIFICOS ATLANTICO FRA-2026-015 CIF B36987123"},
    {"date": "261015", "val": "261015", "cents": 89000,  "doc": "0000000005", "ref1": "TRF BODEGAS LA RIOJA    ", "c2": "TRANSF BODEGAS LA RIOJA FRA-2026-018 CIF B26543210"},
]
tot_abonos = sum(c["cents"] for c in cobros_n43)

for c in cobros_n43:
    # Registro 22 (80 caracteres exactos)
    l22 = "22" + "0049" + "1500" + c["date"] + c["val"] + "04" + "000" + "2" + f"{c['cents']:014d}" + c["doc"].ljust(10) + c["ref1"][:12].ljust(12) + " "*16
    lines_n43.append(l22)
    # Registro 23 (80 caracteres exactos)
    l23 = "23" + "01" + c["c2"].ljust(76)
    lines_n43.append(l23)

# Fin de cuenta 33 (80 caracteres exactos)
final_bal = 1500000 + tot_abonos
l33 = "33" + "0049" + "1500" + "0012345678" + "00000" + "00000000000000" + f"{len(cobros_n43):05d}" + f"{tot_abonos:014d}" + "2" + f"{final_bal:014d}" + "978" + "    "
lines_n43.append(l33)

# Fin de fichero 88 (80 caracteres exactos)
l88 = "88" + "999999" + "000001" + f"{len(lines_n43)+1:06d}" + " "*60
lines_n43.append(l88)

with open(n43_path, "w", encoding="windows-1252", newline="\r\n") as f:
    f.write("\r\n".join(lines_n43) + "\r\n")
print(f"Norma 43 guardado en: {n43_path}")

# ==============================================================================
# 5. FICHERO STRIPE PAYOUT BATCH (.csv)
# ==============================================================================
# Agrupa las facturas FRA-2026-001 (1250€), FRA-2026-002 (890.50€), FRA-2026-004 (695€),
# FRA-2026-012 (1635.50€), FRA-2026-016 (1120.75€) y FRA-2026-019 (1340€)
# Bruto: 6.931,75 € | Comisión Stripe: 98,42 € | Neto Payout: 6.833,33 €
stripe_path = TARGET_DIR / "extracto_stripe_ventas.csv"
stripe_rows = [
    ["id", "Description", "Created (UTC)", "Amount", "Fee", "Net", "Customer Name", "Customer Email", "Invoice / Order"],
    ["py_1PqA01", "Pago pedido tienda online", "2026-10-02 10:15", "1250.00", "17.75", "1232.25", "Logística Exprés Ibérica S.L.", "admin@logisticaexpres.es", "FRA-2026-001"],
    ["py_1PqA02", "Suscripción Cloud Mensual", "2026-10-03 11:30", "890.50", "12.65", "877.85", "Telefónica Soluciones Digitales S.A.", "facturacion@telefonica.es", "FRA-2026-002"],
    ["py_1PqA04", "Embalajes isotérmicos lote", "2026-10-05 09:20", "695.00", "9.87", "685.13", "Distribuciones Alimentarias del Sur S.L.", "compras@distsur.es", "FRA-2026-004"],
    ["py_1PqA12", "Envases vidrio laboratorio", "2026-10-10 16:40", "1635.50", "23.22", "1612.28", "Farmacéutica & Cosmética Cataluña S.A.", "admin@farmacosmetica.cat", "FRA-2026-012"],
    ["py_1PqA16", "Rollos tejido transpirable", "2026-10-12 14:10", "1120.75", "15.91", "1104.84", "Textil y Confección Galicia S.L.", "pedidos@textilgalicia.com", "FRA-2026-016"],
    ["py_1PqA19", "Bolsas envasado al vacío", "2026-10-13 18:05", "1340.00", "19.02", "1320.98", "Embutidos & Jamones de Burgos S.L.", "jamones@embutidosburgos.es", "FRA-2026-019"],
]
with open(stripe_path, "w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f)
    writer.writerows(stripe_rows)
print(f"Stripe Payout guardado en: {stripe_path}")

# ==============================================================================
# 6. GENERACIÓN DE 20 FACTURAS INDIVIDUALES EN PDF (ReportLab)
# ==============================================================================
styles = getSampleStyleSheet()
title_style = ParagraphStyle(
    "TitleStyle",
    parent=styles["Normal"],
    fontName="Helvetica-Bold",
    fontSize=18,
    leading=22,
    textColor=colors.HexColor("#09090b")
)
header_bold = ParagraphStyle(
    "HeaderBold",
    parent=styles["Normal"],
    fontName="Helvetica-Bold",
    fontSize=9,
    leading=12,
    textColor=colors.HexColor("#27272a")
)
body_style = ParagraphStyle(
    "BodyStyle",
    parent=styles["Normal"],
    fontName="Helvetica",
    fontSize=8.5,
    leading=11.5,
    textColor=colors.HexColor("#3f3f46")
)
mono_style = ParagraphStyle(
    "MonoStyle",
    parent=styles["Normal"],
    fontName="Courier",
    fontSize=8.5,
    leading=11,
    textColor=colors.HexColor("#18181b")
)
total_style = ParagraphStyle(
    "TotalStyle",
    parent=styles["Normal"],
    fontName="Helvetica-Bold",
    fontSize=11,
    leading=14,
    textColor=colors.HexColor("#09090b"),
    alignment=2
)

import unicodedata

def clean_filename(s: str) -> str:
    nfkd = unicodedata.normalize('NFKD', s)
    ascii_str = ''.join([c for c in nfkd if not unicodedata.combining(c)])
    clean = re.sub(r'[^a-zA-Z0-9_\-]', '_', ascii_str)
    return re.sub(r'_+', '_', clean).strip('_')

for inv in INVOICES:
    safe_client = clean_filename(inv['client'][:25])
    pdf_filename = PDF_DIR / f"{inv['num']}_{safe_client}.pdf"
    doc = SimpleDocTemplate(
        str(pdf_filename),
        pagesize=A4,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm
    )
    story = []

    # Cabecera Emisor vs Factura
    emisor_info = [
        Paragraph("<b>FINDAFINITY DEMO TECH S.L.</b>", header_bold),
        Paragraph("CIF: B-88776655", body_style),
        Paragraph("Paseo de la Castellana 95, Planta 14", body_style),
        Paragraph("28046 Madrid &bull; España", body_style),
        Paragraph("Email: facturacion@findaffinity.com", body_style),
    ]
    factura_info = [
        Paragraph("FACTURA ORDINARIA", title_style),
        Paragraph(f"<b>Nº Factura:</b> {inv['num']}", header_bold),
        Paragraph(f"<b>Fecha Emisión:</b> {inv['date']}", body_style),
        Paragraph(f"<b>Fecha Vto:</b> {inv['due']}", body_style),
        Paragraph(f"<b>Subcuenta PGC:</b> {inv['subcuenta']}", mono_style),
    ]
    header_table = Table([[emisor_info, factura_info]], colWidths=[9.5 * cm, 8.5 * cm])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 0.6 * cm))

    # Caja del Cliente
    cliente_box = [
        [Paragraph("<b>DATOS FISCALES DEL CLIENTE</b>", header_bold)],
        [Paragraph(f"<b>Razón Social:</b> {inv['client']}", body_style)],
        [Paragraph(f"<b>NIF / CIF:</b> {inv['cif']}", mono_style)],
        [Paragraph(f"<b>Dirección:</b> {inv['address']}", body_style)],
        [Paragraph(f"<b>Población:</b> {inv['city']}", body_style)],
    ]
    t_cliente = Table(cliente_box, colWidths=[18 * cm])
    t_cliente.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f4f4f5")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#e4e4e7")),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_cliente)
    story.append(Spacer(1, 0.6 * cm))

    # Tabla de Líneas
    line_headers = [
        Paragraph("<b>Descripción / Concepto</b>", header_bold),
        Paragraph("<b>Cant.</b>", header_bold),
        Paragraph("<b>Precio Unit.</b>", header_bold),
        Paragraph("<b>IVA</b>", header_bold),
        Paragraph("<b>Total (€)</b>", header_bold),
    ]
    unit_price = inv['base'] / inv['qty']
    line_row = [
        Paragraph(inv['concept'], body_style),
        Paragraph(str(inv['qty']), body_style),
        Paragraph(f"{unit_price:.2f} €", body_style),
        Paragraph(f"{inv['iva_rate']}%", body_style),
        Paragraph(f"{inv['base']:.2f} €", body_style),
    ]
    t_lines = Table([line_headers, line_row], colWidths=[9 * cm, 1.5 * cm, 2.5 * cm, 2 * cm, 3 * cm])
    t_lines.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#18181b")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (1,0), (-1,-1), 'RIGHT'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor("#e4e4e7")),
    ]))
    story.append(t_lines)
    story.append(Spacer(1, 0.6 * cm))

    # Resumen de Totales
    totals_data = [
        ["Base Imponible:", f"{inv['base']:.2f} €"],
        [f"Cuota IVA ({inv['iva_rate']}%):", f"{inv['iva']:.2f} €"],
        ["TOTAL FACTURA:", f"{inv['total']:.2f} €"],
    ]
    t_totals = Table(totals_data, colWidths=[14 * cm, 4 * cm])
    t_totals.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'RIGHT'),
        ('FONTNAME', (0,0), (-1,-2), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-2), 9),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,-1), (-1,-1), 11),
        ('LINEABOVE', (0,-1), (-1,-1), 1, colors.HexColor("#18181b")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_totals)
    story.append(Spacer(1, 0.8 * cm))

    # Información de Pago y Pie
    pago_info = [
        Paragraph(f"<b>Forma de Pago:</b> Transferencia bancaria a <b>ES91 2100 0418 4502 0005 1234</b> (CaixaBank)", body_style),
        Paragraph(f"<b>Referencia obligatoria:</b> {inv['num']} / {inv['cif']}", body_style),
        Paragraph("<i>Inscrita en el Registro Mercantil de Madrid, Tomo 34.567, Folio 120, Hoja M-621.890.</i>", ParagraphStyle("Footer", parent=body_style, fontSize=7, textColor=colors.HexColor("#71717a"))),
    ]
    for p in pago_info:
        story.append(p)
        story.append(Spacer(1, 2))

    doc.build(story)

print(f"20 Facturas en PDF creadas en: {PDF_DIR}")

# ==============================================================================
# 7. DOCUMENTO GUÍA DE PRUEBAS
# ==============================================================================
readme_content = f"""# Carpeta de Pruebas: Facturas y Extractos para FINDAFINITY

Esta carpeta contiene un juego completo de **20 facturas comerciales inventadas** y **ficheros de extracto bancario/pasarela** listos para ser probados de inmediato en la plataforma FINDAFINITY.

---

## 📁 Contenido de la Carpeta

| Fichero | Formato | Descripción |
| :--- | :--- | :--- |
| **`cartera_facturas_emitidas_20.xlsx`** | Excel (.xlsx) | Cartera viva con 20 facturas, importes, CIFs, subcuentas PGC 430 y formato oficial español. |
| **`cartera_facturas_emitidas_20.csv`** | CSV (.csv) | Cartera de 20 facturas con delimitador `;` y codificación UTF-8 para importación ERP. |
| **`cartera_facturas_emitidas_20.json`** | JSON (.json) | Cartera en JSON estructurado para integración por API. |
| **`extracto_bancario_norma43.n43`** | Norma 43 (.n43) | Extracto bancario real con 5 transferencias coincidentes por importe, CIF y número de factura. |
| **`extracto_stripe_ventas.csv`** | Stripe CSV (.csv) | Liquidación neta agrupada con 6 facturas y cálculo de comisiones (cuenta 626). |
| **`facturas_pdf/`** | 20 PDFs individuales | Facturas imprimibles maquetadas profesionalmente con desglose de Base, IVA 21%, vencimiento y datos fiscales. |

---

## 🚀 Cómo probarlas en FINDAFINITY (Paso a Paso)

### Escenario A: Conciliación Bancaria B2B (Norma 43 + Cartera)
1. Inicia sesión en FINDAFINITY (o abre el Radar de Conciliación).
2. Asegúrate de tener seleccionada tu **Empresa Cliente** en la cabecera.
3. En la **Zona de Ingesta**, arrastra o selecciona:
   - **Extracto Bancario:** Selecciona `extracto_bancario_norma43.n43`.
   - **Cartera de Facturas (Opcional pero recomendado):** Selecciona `cartera_facturas_emitidas_20.xlsx` (o `.csv`).
4. Haz clic en **Procesar Extracto**.
5. **Resultado esperado:**
   - 5 partidas bancarias conciliadas al 100% de confianza (**Semáforo VERDE**).
   - Coincidencia perfecta de NIF/CIF y número de documento (FRA-2026-003, FRA-2026-008, FRA-2026-009, FRA-2026-015, FRA-2026-018).
   - Asientos de partida doble PGC generados con **Descuadre 0,00 €**:
     - `(D) 572000001 Bancos`
     - `(H) 43000000x Clientes`

### Escenario B: Conciliación de Pasarela Stripe con Comisiones
1. En la **Zona de Ingesta**, sube:
   - **Extracto Bancario / Pasarela:** Selecciona `extracto_stripe_ventas.csv`.
   - **Cartera de Facturas:** Selecciona `cartera_facturas_emitidas_20.xlsx`.
2. Haz clic en **Procesar Extracto**.
3. **Resultado esperado:**
   - Se reconocen las 6 transacciones de Stripe.
   - Desglose contable automático a 3 patas:
     - `(D) 572 Bancos (Importe Neto cobrado)`
     - `(D) 626 Comisiones de pasarela Stripe`
     - `(H) 430 Clientes (Importe Bruto de la factura)`
   - Descuadre Debe - Haber: **0,00 €**.

---

## 📊 Resumen de las 20 Facturas

- **Total Cartera:** {total_suma:,.2f} €
- **Número de facturas:** 20
- **Empresas deudoras:** Logística Exprés, Telefónica Soluciones, Almacenes Centrales, Novatech, Amazon EU, Bodegas Rioja, etc.
- **Rango de importes:** Desde 420,00 € hasta 5.600,00 €.
"""

readme_path = TARGET_DIR / "README_INSTRUCCIONES.md"
with open(readme_path, "w", encoding="utf-8") as f:
    f.write(readme_content)
print(f"Instrucciones guardadas en: {readme_path}")
