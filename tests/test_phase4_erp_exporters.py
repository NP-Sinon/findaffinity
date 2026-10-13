"""
Test Suite para la Fase 4: Exportadores Contables Nativos para ERPs Españoles
Verifica:
1. Exportador Wolters Kluwer A3 (protocolo Suenlace / a3conv).
2. Exportador Sage 50 / Sage Despachos (estándar XDIARIO.csv).
3. Exportador DELSOL Contasol (hoja de cálculo XLSX nativa).
4. Rechazo preventivo de asientos contables no balanceados en todos los exportadores.
"""
from datetime import date
import io
import openpyxl
import pytest

from src.accounting.double_entry import AccountingEntryBuilder, DoubleEntryJournal, JournalEntryLine
from src.accounting.exporters.a3_suenlace import A3SuenlaceExporter
from src.accounting.exporters.sage_xdiario import SageXdiarioExporter
from src.accounting.exporters.contasol_xlsx import ContasolXlsxExporter


@pytest.fixture
def sample_balanced_journal() -> DoubleEntryJournal:
    """Crea un asiento de liquidación perfectamente balanceado a cero."""
    # 572 (Bancos): 98.00 € (9800 cents) al Debe
    # 626 (Comisiones): 2.00 € (200 cents) al Debe
    # 430 (Cliente): 100.00 € (10000 cents) al Haber
    journal = DoubleEntryJournal(
        entry_number=42,
        entry_date=date(2026, 10, 15),
        description="Liquidación TPV Redsys 15/10/2026",
    )
    journal.lines = [
        JournalEntryLine(
            account_code="572000001",
            account_name="Banco Santander",
            concept="Cobro Liq Redsys",
            document_number="LIQ-20261015",
            debit_cents=9800,
            credit_cents=0,
            contrapartida="626000001",
        ),
        JournalEntryLine(
            account_code="626000001",
            account_name="Comisiones bancarias",
            concept="Comisión TPV Redsys",
            document_number="FEE-20261015",
            debit_cents=200,
            credit_cents=0,
            contrapartida="572000001",
        ),
        JournalEntryLine(
            account_code="430000123",
            account_name="Cliente Comercio S.L.",
            concept="Fra. F-2026-99",
            document_number="F-2026-99",
            debit_cents=0,
            credit_cents=10000,
            contrapartida="572000001",
        ),
    ]
    journal.check_balance()
    return journal


def test_a3_suenlace_exporter(sample_balanced_journal):
    """Verifica la exportación al formato de texto delimitado Suenlace de Wolters Kluwer A3."""
    output_text = A3SuenlaceExporter.export_journal_entries([sample_balanced_journal])
    lines = [l for l in output_text.splitlines() if l.strip()]

    assert len(lines) == 3
    # Formato: Asiento;Fecha;Cuenta;Signo;Importe;Concepto;Documento;Contrapartida
    # Línea 1: Banco Santander
    parts_1 = lines[0].split(";")
    assert parts_1[0] == "000042"          # Asiento formateado
    assert parts_1[1] == "15102026"        # Fecha DDMMAAAA
    assert parts_1[2] == "572000001"       # Cuenta 9 dígitos
    assert parts_1[3] == "D"               # Signo Debe
    assert parts_1[4] == "98.00"           # Importe 98.00
    assert "Cobro Liq Redsys" in parts_1[5]

    # Línea 3: Cliente (Haber)
    parts_3 = lines[2].split(";")
    assert parts_3[2] == "430000123"
    assert parts_3[3] == "H"
    assert parts_3[4] == "100.00"


def test_sage_xdiario_exporter(sample_balanced_journal):
    """Verifica la exportación a XDIARIO.csv para Sage 50 y Sage Despachos."""
    csv_text = SageXdiarioExporter.export_xdiario_csv([sample_balanced_journal], delimiter=";")
    lines = [l for l in csv_text.splitlines() if l.strip()]

    # Cabecera + 3 líneas
    assert len(lines) == 4
    header = lines[0].split(";")
    assert header == ["Asiento", "Fecha", "Cuenta", "Concepto", "Documento", "Debe", "Haber", "Factura", "Contrapartida"]

    # Fila Debe banco
    row_1 = lines[1].split(";")
    assert row_1[0] == "42"
    assert row_1[1] == "15/10/2026"
    assert row_1[2] == "572000001"
    assert row_1[5] == "98.00"
    assert row_1[6] == "0.00"

    # Fila Haber cliente
    row_3 = lines[3].split(";")
    assert row_3[2] == "430000123"
    assert row_3[5] == "0.00"
    assert row_3[6] == "100.00"


def test_contasol_xlsx_exporter(sample_balanced_journal):
    """Verifica la exportación binaria a Excel XLSX para DELSOL Contasol."""
    xlsx_bytes = ContasolXlsxExporter.export_xlsx([sample_balanced_journal])
    assert len(xlsx_bytes) > 0

    # Cargar el libro desde memoria con openpyxl para verificar celdas
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title == "Asientos Contables"

    # Fila 1: Cabeceras
    headers = [cell.value for cell in ws[1]]
    assert headers == ["Diario", "Fecha", "Asiento", "Subcuenta", "Concepto", "Documento", "Debe", "Haber", "Contrapartida"]

    # Filas de datos
    assert ws.max_row == 4
    # Fila 2: Banco
    row_2 = [cell.value for cell in ws[2]]
    assert row_2[0] == 1               # Diario
    assert row_2[1] == "15/10/2026"    # Fecha
    assert row_2[2] == 42              # Asiento
    assert row_2[3] == "572000001"     # Subcuenta
    assert row_2[6] == 98.00           # Debe
    assert row_2[7] == 0.00            # Haber

    # Fila 4: Cliente
    row_4 = [cell.value for cell in ws[4]]
    assert row_4[3] == "430000123"
    assert row_4[6] == 0.00
    assert row_4[7] == 100.00


def test_exporters_reject_unbalanced_journals():
    """Verifica que todos los exportadores bloquean la emisión si un asiento contable no suma cero."""
    unbalanced = DoubleEntryJournal(
        entry_number=99,
        entry_date=date(2026, 10, 15),
        description="Asiento corrupto",
    )
    unbalanced.lines = [
        JournalEntryLine("572000001", "Banco", "Test", "DOC1", debit_cents=10000, credit_cents=0),
        JournalEntryLine("430000001", "Cliente", "Test", "DOC2", debit_cents=0, credit_cents=9500), # Descuadre de 500 cts
    ]

    with pytest.raises(ValueError, match="descuadrado|no está cuadrado"):
        A3SuenlaceExporter.export_journal_entries([unbalanced])

    with pytest.raises(ValueError, match="descuadre"):
        SageXdiarioExporter.export_xdiario_csv([unbalanced])

    with pytest.raises(ValueError, match="descuadrado"):
        ContasolXlsxExporter.export_xlsx([unbalanced])
